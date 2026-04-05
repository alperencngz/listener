"""DOCX (Word) export for meeting transcripts."""

import io
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def export_docx(
    session_id: str,
    title: str,
    date_str: str,
    duration_str: str,
    language: str,
    transcript_text: str,
    analysis_text: str | None = None,
) -> io.BytesIO:
    """Generate a DOCX file and return it as a BytesIO stream.

    Args:
        session_id: Session identifier (e.g., "2026-04-04_14-30-00").
        title: Meeting title.
        date_str: Human-readable date string.
        duration_str: Duration (e.g., "45m 12s").
        language: Language code.
        transcript_text: Full timestamped transcript markdown text.
        analysis_text: Optional analysis markdown text.

    Returns:
        BytesIO containing the DOCX file bytes.
    """
    doc = Document()

    # Title
    heading = doc.add_heading(title, level=1)

    # Metadata paragraph
    meta = doc.add_paragraph()
    meta.add_run(f"Date: {date_str}").bold = True
    meta.add_run(f"  |  Duration: {duration_str}  |  Language: {language.upper()}")
    meta.space_after = Pt(12)

    doc.add_paragraph("\u2500" * 60)

    # Transcript section
    doc.add_heading("Transcript", level=2)
    _add_transcript_to_docx(doc, transcript_text)

    # Analysis section (if available)
    if analysis_text:
        doc.add_page_break()
        doc.add_heading("Analysis", level=2)
        _add_markdown_to_docx(doc, analysis_text)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def _add_transcript_to_docx(doc: Document, transcript_text: str):
    """Add transcript lines with styled timestamps."""
    lines = transcript_text.strip().split("\n")
    for line in lines:
        line = line.strip()
        if not line:
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)

        # Match timestamps: [HH:MM:SS] or **[HH:MM:SS] SPEAKER:** patterns
        ts_match = re.match(
            r'(?:\*\*)?\[(\d{1,2}:\d{2}(?::\d{2})?)\](?:\s*([A-Z_0-9]+):?\*\*)?(.*)$',
            line,
        )
        if ts_match:
            ts_val = ts_match.group(1)
            speaker = ts_match.group(2) or ""
            text = ts_match.group(3).strip()

            # Timestamp in blue Courier New
            ts_run = p.add_run(f"[{ts_val}] ")
            ts_run.font.name = "Courier New"
            ts_run.font.size = Pt(10)
            ts_run.font.color.rgb = RGBColor(0x63, 0x66, 0xF1)

            # Speaker label if present
            if speaker:
                spk_run = p.add_run(f"{speaker}: ")
                spk_run.font.bold = True
                spk_run.font.size = Pt(11)

            # Text
            text_run = p.add_run(text)
            text_run.font.size = Pt(11)
        else:
            # Plain text line
            run = p.add_run(line)
            run.font.size = Pt(11)


def _add_markdown_to_docx(doc: Document, markdown_text: str):
    """Convert basic markdown to Word styles."""
    lines = markdown_text.strip().split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        # Skip empty lines
        if not line:
            i += 1
            continue

        # Headings
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("---"):
            doc.add_paragraph("\u2500" * 60)
        elif line.startswith("- ") or line.startswith("* "):
            # Bullet list
            p = doc.add_paragraph(line[2:].strip(), style="List Bullet")
        elif re.match(r'^\d+\.\s', line):
            # Numbered list
            text = re.sub(r'^\d+\.\s*', '', line)
            p = doc.add_paragraph(text, style="List Number")
        else:
            # Regular paragraph — strip bold markers
            clean = line.replace("**", "")
            doc.add_paragraph(clean)

        i += 1
